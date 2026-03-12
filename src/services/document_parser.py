"""Document parsing service supporting PDF, DOCX, OCR, and legacy DOC fallbacks."""

import logging
import os
import re
import subprocess
import tempfile
from io import BytesIO
from pathlib import Path
from typing import Callable, List, Literal, Tuple

from PIL import Image
from src.core.config import settings

logger = logging.getLogger(__name__)

WATERMARK_TOKEN_PATTERN = re.compile(r"[0-9a-f]{16}HR-[A-Za-z0-9~_-]{8,}")
NORMAL_CHINESE_PUNCT_PATTERN = re.compile(
    r"[\uFF0C\u3002\uFF1B\uFF1A\u3001\uFF08\uFF09\u300A\u300B\u3010\u3011]"
)
RESUME_KEYWORDS = [
    "\u59d3\u540d",
    "\u7535\u8bdd",
    "\u90ae\u7bb1",
    "\u5de5\u4f5c",
    "\u7ecf\u5386",
    "\u9879\u76ee",
    "\u6559\u80b2",
    "\u6280\u80fd",
    "\u8d1f\u8d23",
    "\u6570\u636e",
    "\u5f00\u53d1",
    "\u5de5\u7a0b\u5e08",
    "\u516c\u53f8",
    "\u5c97\u4f4d",
    "\u804c\u8d23",
    "\u6210\u679c",
]


class DocumentParser:
    """Parse resume files and return plain text."""

    def parse(self, file_path: str) -> str:
        file_kind = self._detect_file_kind(file_path)

        if file_kind == "pdf":
            return self._parse_pdf(file_path)
        if file_kind == "docx":
            return self._parse_docx(file_path)
        if file_kind == "doc":
            converted_path = self._convert_doc_to_pdf(file_path)
            try:
                return self._parse_pdf(converted_path)
            finally:
                converted = Path(converted_path)
                converted.unlink(missing_ok=True)
                if converted.parent.name.startswith("resume-doc-convert-"):
                    converted.parent.rmdir()
        if file_kind == "image":
            return self._extract_image_with_ocr(file_path)
        raise ValueError(f"Unsupported file type: {Path(file_path).suffix.lower()}")

    def _detect_file_kind(self, file_path: str) -> Literal["pdf", "docx", "doc", "image", "unknown"]:
        path = Path(file_path)
        suffix = path.suffix.lower()
        header = path.read_bytes()[:16]

        if header.startswith(b"%PDF"):
            return "pdf"
        if header.startswith(b"PK"):
            return "docx"
        if header.startswith(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"):
            return "doc"
        if header.startswith((b"\x89PNG\r\n\x1a\n", b"\xff\xd8\xff")):
            return "image"

        if suffix == ".pdf":
            return "pdf"
        if suffix == ".doc":
            return "doc"
        if suffix == ".docx":
            return "docx"
        if suffix in {".png", ".jpg", ".jpeg"}:
            return "image"
        return "unknown"

    def _parse_pdf(self, file_path: str) -> str:
        try:
            extractors: List[Tuple[str, Callable[[str], str]]] = [
                ("pypdf2", self._extract_pdf_with_pypdf2),
            ]

            if self._has_module("fitz"):
                extractors.append(("pymupdf", self._extract_pdf_with_pymupdf))
            if self._has_module("pdfplumber"):
                extractors.append(("pdfplumber", self._extract_pdf_with_pdfplumber))

            best_text = ""
            best_score = -1

            for name, extractor in extractors:
                try:
                    raw_text = extractor(file_path)
                except Exception as extractor_error:
                    logger.warning("PDF extractor %s failed: %s", name, extractor_error)
                    continue

                cleaned_text = self._clean_text(raw_text)
                if not cleaned_text:
                    continue

                score = self._text_quality_score(cleaned_text)
                if score > best_score:
                    best_score = score
                    best_text = cleaned_text

                if not self._is_low_quality_text(cleaned_text):
                    logger.info(
                        "Parsed PDF with %s: %s, length: %s chars",
                        name,
                        file_path,
                        len(cleaned_text),
                    )
                    return cleaned_text

                logger.warning(
                    "Low quality PDF text from %s (len=%s, score=%s), trying fallback",
                    name,
                    len(cleaned_text),
                    score,
                )

            if best_text:
                logger.warning(
                    "Parsed PDF with degraded quality after trying fallbacks: %s, length=%s",
                    file_path,
                    len(best_text),
                )
                ocr_text = self._extract_pdf_with_ocr(file_path)
                if ocr_text and len(ocr_text) > len(best_text):
                    return ocr_text
                return best_text

            ocr_text = self._extract_pdf_with_ocr(file_path)
            if ocr_text:
                logger.info(
                    "Parsed PDF via OCR fallback: %s, length: %s chars",
                    file_path,
                    len(ocr_text),
                )
                return ocr_text

            raise ValueError("Failed to extract readable text from PDF")
        except Exception as e:
            logger.error("Failed to parse PDF: %s, error: %s", file_path, e)
            raise

    def _extract_pdf_with_pypdf2(self, file_path: str) -> str:
        import PyPDF2

        text_parts = []
        with open(file_path, "rb") as file:
            pdf_reader = PyPDF2.PdfReader(file)
            for page in pdf_reader.pages:
                page_text = page.extract_text()
                if page_text:
                    text_parts.append(page_text)
        return "\n".join(text_parts)

    def _extract_pdf_with_pymupdf(self, file_path: str) -> str:
        import fitz

        text_parts = []
        doc = fitz.open(file_path)
        try:
            for page in doc:
                page_text = page.get_text("text")
                if page_text:
                    text_parts.append(page_text)
        finally:
            doc.close()
        return "\n".join(text_parts)

    def _extract_pdf_with_pdfplumber(self, file_path: str) -> str:
        import pdfplumber

        text_parts = []
        with pdfplumber.open(file_path) as pdf:
            for page in pdf.pages:
                page_text = page.extract_text() or ""
                if page_text:
                    text_parts.append(page_text)
        return "\n".join(text_parts)

    def _parse_docx(self, file_path: str) -> str:
        try:
            from docx import Document

            doc = Document(file_path)
            paragraphs = [para.text for para in doc.paragraphs if para.text.strip()]

            tables_text = []
            for table in doc.tables:
                for row in table.rows:
                    row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_text:
                        tables_text.append(" | ".join(row_text))

            text = "\n".join(paragraphs + tables_text)
            text = self._clean_text(text)

            logger.info("Parsed DOCX: %s, length: %s chars", file_path, len(text))
            return text
        except Exception as e:
            logger.error("Failed to parse DOCX: %s, error: %s", file_path, e)
            raise

    def _extract_pdf_with_ocr(self, file_path: str) -> str:
        if not self._has_module("rapidocr_onnxruntime"):
            raise RuntimeError("RapidOCR is not installed")

        import fitz

        text_parts = []
        doc = fitz.open(file_path)
        try:
            for page in doc:
                pix = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
                image = Image.open(BytesIO(pix.tobytes("png")))
                try:
                    page_text = self._run_ocr(image)
                finally:
                    image.close()
                if page_text:
                    text_parts.append(page_text)
        finally:
            doc.close()

        return self._clean_text("\n".join(text_parts))

    def _extract_image_with_ocr(self, file_path: str) -> str:
        if not self._has_module("rapidocr_onnxruntime"):
            raise RuntimeError("RapidOCR is not installed")

        image = Image.open(file_path)
        try:
            return self._clean_text(self._run_ocr(image))
        finally:
            image.close()

    def _run_ocr(self, image: Image.Image) -> str:
        from rapidocr_onnxruntime import RapidOCR

        engine = RapidOCR()
        result, _ = engine(image)
        if not result:
            return ""
        return "\n".join(item[1] for item in result if len(item) >= 2 and item[1])

    def _convert_doc_to_pdf(self, file_path: str) -> str:
        soffice_path = settings.libreoffice_path or os.getenv(
            "LIBREOFFICE_PATH",
            r"C:\Program Files\LibreOffice\program\soffice.exe",
        )
        output_dir = Path(tempfile.mkdtemp(prefix="resume-doc-convert-"))
        command = [
            soffice_path,
            "--headless",
            "--convert-to",
            "pdf",
            "--outdir",
            str(output_dir),
            file_path,
        ]
        try:
            subprocess.run(command, check=True, capture_output=True, text=True)
        except FileNotFoundError as exc:
            raise RuntimeError("LibreOffice is not installed") from exc
        except subprocess.CalledProcessError as exc:
            raise RuntimeError(
                f"LibreOffice doc conversion failed: {exc.stderr or exc.stdout}"
            ) from exc

        converted_path = output_dir / f"{Path(file_path).stem}.pdf"
        if not converted_path.exists():
            raise RuntimeError("LibreOffice doc conversion did not produce a PDF")
        return str(converted_path)

    def _clean_text(self, text: str) -> str:
        if not text:
            return ""

        text = self._repair_text_encoding(text)

        text = WATERMARK_TOKEN_PATTERN.sub("", text)

        text = re.sub(r"\n{3,}", "\n\n", text)
        text = re.sub(r" {2,}", " ", text)

        lines = [line.strip() for line in text.splitlines()]
        deduped_lines = []
        for line in lines:
            if not line:
                deduped_lines.append("")
                continue

            line = self._normalize_line(line)
            if not line:
                continue
            if self._is_noise_line(line):
                continue
            if deduped_lines and deduped_lines[-1] == line:
                continue
            deduped_lines.append(line)

        return "\n".join(deduped_lines).strip()

    def _normalize_line(self, line: str) -> str:
        if not line:
            return ""

        has_chinese = bool(re.search(r"[\u4e00-\u9fff]", line))
        if not has_chinese:
            return line.strip()

        line = re.sub(r"\s+[A-Za-z](?=\s|$)", "", line)
        line = re.sub(r"\s+[A-Za-z](?=[\u4e00-\u9fff])", " ", line)
        line = re.sub(r"^\d+\s+\d+\s+(?=[\u4e00-\u9fff])", "", line)
        line = re.sub(r"\s+\d\s+\d$", "", line)

        return line.strip()

    def _is_noise_line(self, line: str) -> bool:
        if not line:
            return False

        if WATERMARK_TOKEN_PATTERN.search(line):
            return True

        if re.fullmatch(r"^[~`!@#$%^&*()_+\-=\[\]{};,.<>/?\\|:]+$", line):
            return True

        has_chinese = bool(re.search(r"[\u4e00-\u9fff]", line))
        if not has_chinese:
            compact = re.sub(r"[^A-Za-z0-9]", "", line)

            if len(compact) <= 2:
                return True

            if (
                "@" not in line
                and re.fullmatch(r"[A-Za-z0-9~_-]{20,}", line)
                and sum(ch.isdigit() for ch in line) >= 4
            ):
                return True

            if re.fullmatch(r"(?=.*[A-Z])(?=.*[a-z])(?=.*\d)[A-Za-z0-9]{10,}", compact):
                return True

        return False

    def _is_low_quality_text(self, text: str) -> bool:
        if not text:
            return True

        noise_hits = len(WATERMARK_TOKEN_PATTERN.findall(text))
        if noise_hits >= 3:
            return True

        lines = [line.strip() for line in text.splitlines() if line.strip()]
        chinese_chars = len(re.findall(r"[\u4e00-\u9fff]", text))
        long_meaningful_lines = sum(1 for line in lines if len(line) >= 12)

        if chinese_chars < 120:
            return True
        if long_meaningful_lines < 5:
            return True
        return False

    def _text_quality_score(self, text: str) -> int:
        chinese_chars = len(re.findall(r"[\u4e00-\u9fff]", text))
        lines = [line.strip() for line in text.splitlines() if line.strip()]
        long_lines = sum(1 for line in lines if len(line) >= 12)
        noise_hits = len(WATERMARK_TOKEN_PATTERN.findall(text))
        normal_punct = len(NORMAL_CHINESE_PUNCT_PATTERN.findall(text))
        keyword_hits = sum(text.count(keyword) for keyword in RESUME_KEYWORDS)

        return (
            chinese_chars
            + (long_lines * 8)
            + (normal_punct * 12)
            + (keyword_hits * 100)
            - (noise_hits * 200)
        )

    def _repair_text_encoding(self, text: str) -> str:
        base_score = self._encoding_quality_score(text)
        best_text = text
        best_score = base_score

        for source_encoding in ("gbk", "gb18030", "latin1"):
            try:
                repaired = text.encode(source_encoding, errors="ignore").decode(
                    "utf-8", errors="ignore"
                )
            except Exception:
                continue

            if not repaired:
                continue

            score = self._encoding_quality_score(repaired)
            if score > best_score:
                best_score = score
                best_text = repaired

        if best_text != text and best_score >= base_score + 20:
            logger.info(
                "Applied encoding repair for extracted text (score %s -> %s)",
                base_score,
                best_score,
            )
            return best_text
        return text

    def _encoding_quality_score(self, text: str) -> int:
        chinese_chars = len(re.findall(r"[\u4e00-\u9fff]", text))
        normal_punct = len(NORMAL_CHINESE_PUNCT_PATTERN.findall(text))
        keyword_hits = sum(text.count(keyword) for keyword in RESUME_KEYWORDS)
        replacement_chars = text.count("�")
        return (
            chinese_chars
            + (normal_punct * 12)
            + (keyword_hits * 100)
            - (replacement_chars * 40)
        )

    def _has_module(self, module_name: str) -> bool:
        try:
            __import__(module_name)
            return True
        except ImportError:
            return False

    def get_file_info(self, file_path: str) -> dict:
        path = Path(file_path)
        return {
            "filename": path.name,
            "extension": path.suffix.lower(),
            "size_bytes": path.stat().st_size if path.exists() else 0,
        }
